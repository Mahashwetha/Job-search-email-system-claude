"""
CV Tailor - evidence-only edits to the current CV DOCX for one job.

The comparison (which JD skills are missing, which ones the candidate confirms) is done by
the resume-tailor skill in Claude Code; this script only does the deterministic parts.

Usage:
  python cv_tailor.py jd <job-url>                 print the job description text
  python cv_tailor.py cv                           print the base CV, one paragraph per line
  python cv_tailor.py apply "<Company>" <edits.json>

edits.json:
  {"replace": [{"old": "exact text inside one run", "new": "replacement"}, ...],
   "add_skills": ["Cucumber (TDD)", "YourKit Java Profiler"]}

Output: RESUME_OUTPUT_DIR/resume_<company>.docx + .pdf (PDF via Microsoft Word).
"""

import html
import json
import os
import re
import sys

import requests
from bs4 import BeautifulSoup
from docx import Document

try:
    import config
except ImportError:
    print("ERROR: config.py not found!")
    sys.exit(1)

BASE_CV_DOCX = getattr(config, "BASE_CV_DOCX", os.path.join(os.path.expanduser("~"), "resume.docx"))
SKILLS_LABEL = "Frameworks & Tools:"
NL = "\n"


def _console_utf8():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")


# ── Job description fetching ─────────────────────────────────────────────────

def _html_text(markup):
    return BeautifulSoup(markup or "", "html.parser").get_text(NL, strip=True)


def _ashby_jd(url):
    """Ashby pages are JS-rendered; its public posting API returns the full description."""
    m = re.search(r"jobs\.ashbyhq\.com/([^/?#]+)/([0-9a-f-]{36})", url)
    if not m:
        return ""
    org, job_id = m.groups()
    resp = requests.get(f"https://api.ashbyhq.com/posting-api/job-board/{org}", timeout=20)
    if resp.status_code != 200:
        return ""
    job = next((j for j in resp.json().get("jobs", []) if j.get("id") == job_id), None)
    if not job:
        return ""
    head = f"Job Title: {job.get('title', '')}{NL}Location: {job.get('location', '')}"
    return head + NL + NL + _html_text(job.get("descriptionHtml"))


def _greenhouse_jd(url):
    """job-boards / boards.greenhouse.io pages: the public board API has the full description."""
    m = re.search(r"greenhouse\.io/([^/?#&]+)/jobs/(\d+)", url)
    if not m:
        return ""
    org, job_id = m.groups()
    resp = requests.get(f"https://boards-api.greenhouse.io/v1/boards/{org}/jobs/{job_id}?content=true", timeout=20)
    if resp.status_code != 200:
        return ""
    job = resp.json()
    head = f"Job Title: {job.get('title', '')}{NL}Location: {(job.get('location') or {}).get('name', '')}"
    return head + NL + NL + _html_text(html.unescape(job.get("content", "")))


def _wttj_jd(url):
    """WTTJ pages are JS-rendered; its Algolia index has the summary, missions and profile."""
    m = re.search(r"/companies/([^/]+)/jobs/([^?&#/]+)", url)
    if not m:
        return ""
    job_slug = m.group(2)
    resp = requests.post(
        "https://CSEKHVMS53-dsn.algolia.net/1/indexes/wttj_jobs_production_fr/query",
        headers={"X-Algolia-Application-Id": "CSEKHVMS53", "X-Algolia-API-Key": "4bd8f6215d0cc52b26430765769e65a0",
                 "Origin": "https://www.welcometothejungle.com", "Referer": "https://www.welcometothejungle.com/"},
        json={"params": f"query={job_slug.replace('-', ' ')}&hitsPerPage=10"}, timeout=20)
    if resp.status_code != 200:
        return ""
    hit = next((h for h in resp.json().get("hits", []) if h.get("slug") == job_slug), None)
    if not hit:
        return ""
    missions = NL.join(f"- {x}" for x in hit.get("key_missions") or [])
    parts = [f"Job Title: {hit.get('name', '')}", hit.get("summary") or "",
             ("Key missions:" + NL + missions) if missions else "", hit.get("profile") or ""]
    return (NL + NL).join(x for x in parts if x)


def _site_jd(url):
    """Clean description from the site's own API, or '' if the URL isn't a supported site."""
    import fit_scorer
    u = url.lower()
    if "linkedin.com/jobs/view/" in u:
        return fit_scorer._fetch_linkedin(url)
    if "welcometothejungle.com" in u:
        return _wttj_jd(url)
    if "ashbyhq.com" in u:
        return _ashby_jd(url)
    if "greenhouse.io" in u:
        return _greenhouse_jd(url)
    return ""


def cmd_jd(url):
    """Site API first (LinkedIn, WTTJ, Ashby, Greenhouse); otherwise the longest page-text extraction."""
    import fit_scorer
    import resume_tailor
    text = ""
    try:
        text = _site_jd(url) or ""
    except Exception:
        pass
    if len(text.strip()) < 200:
        texts = []
        for f in (fit_scorer.fetch_job_description, resume_tailor.fetch_job_description):
            try:
                texts.append(f(url) or "")
            except Exception:
                pass
        text = max(texts, key=len, default="")
    if len(text.strip()) < 200:
        print("ERROR: could not fetch the job description. Paste it into the chat instead.")
        sys.exit(1)
    print(text)


# ── CV editing ───────────────────────────────────────────────────────────────

def cmd_cv():
    for i, p in enumerate(Document(BASE_CV_DOCX).paragraphs):
        if p.text.strip():
            print(f"{i}: {p.text}")


def _replace_in_runs(doc, old, new):
    for p in doc.paragraphs:
        for r in p.runs:
            if old in r.text:
                r.text = r.text.replace(old, new, 1)
                return
    raise SystemExit(f"ERROR: text not found inside a single run: {old!r}")


def _add_skills(doc, skills):
    p = next((p for p in doc.paragraphs if p.text.startswith(SKILLS_LABEL)), None)
    if p is None:
        raise SystemExit(f"ERROR: no paragraph starting with {SKILLS_LABEL!r}")
    last = next(r for r in reversed(p.runs) if r.text.strip())
    last.text = last.text.rstrip() + ", " + ", ".join(skills)


def _export_pdf(docx_path):
    import win32com.client
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    try:
        d = word.Documents.Open(docx_path, ReadOnly=True)
        pages = d.ComputeStatistics(2)  # wdStatisticPages
        d.ExportAsFixedFormat(docx_path[:-5] + ".pdf", 17)  # wdExportFormatPDF
        d.Close(False)
    finally:
        word.Quit()
    return pages


def word_diff(base_path, new_path):
    """Changed words per paragraph: [-removed-] {+added+}."""
    import difflib
    base = [p.text for p in Document(base_path).paragraphs]
    new = [p.text for p in Document(new_path).paragraphs]
    for i, (a, b) in enumerate(zip(base, new)):
        if a == b:
            continue
        aw, bw = a.split(), b.split()
        parts = []
        for op, i1, i2, j1, j2 in difflib.SequenceMatcher(None, aw, bw).get_opcodes():
            if op == "equal":
                continue
            ctx = " ".join(aw[max(0, i1 - 2):i1])
            old = " ".join(aw[i1:i2])
            new_ = " ".join(bw[j1:j2])
            parts.append(f"...{ctx} " + (f"[-{old}-] " if old else "") + (f"{{+{new_}+}}" if new_ else ""))
        print(f"Para {i}: " + "  |  ".join(parts))


def cmd_apply(company, edits_path):
    from resume_tailor import safe_company_name
    with open(edits_path, encoding="utf-8") as f:
        edits = json.load(f)
    doc = Document(BASE_CV_DOCX)
    for e in edits.get("replace", []):
        _replace_in_runs(doc, e["old"], e["new"])
    if edits.get("add_skills"):
        _add_skills(doc, edits["add_skills"])

    os.makedirs(config.RESUME_OUTPUT_DIR, exist_ok=True)
    out = os.path.join(config.RESUME_OUTPUT_DIR, f"resume_{safe_company_name(company)}.docx")
    doc.save(out)
    pages = _export_pdf(out)
    print(f"Saved: {out}{NL}Saved: {out[:-5]}.pdf{NL}Pages: {pages}{NL}{NL}Words changed:")
    word_diff(BASE_CV_DOCX, out)


def main():
    _console_utf8()
    args = sys.argv[1:]
    if args[:1] == ["jd"] and len(args) == 2:
        cmd_jd(args[1])
    elif args == ["cv"]:
        cmd_cv()
    elif args[:1] == ["apply"] and len(args) == 3:
        cmd_apply(args[1], args[2])
    else:
        print(__doc__)
        sys.exit(1)


if __name__ == "__main__":
    main()
