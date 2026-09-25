"""
End-to-end test: talks to server.py over stdio exactly like an AI agent would.
Always runs against COPIES, never the real files. The test dir must contain copies of:
  List.xlsx, daily_hot_jobs.json, rejected_remote.json, previous_jobs.json

Run:  tracker_mcp/.venv/Scripts/python.exe tracker_mcp/test_server.py <test-dir>
"""
import asyncio
import json
import os
import sys

from mcp import ClientSession, StdioServerParameters
from mcp.client.stdio import stdio_client

HERE = os.path.dirname(os.path.abspath(__file__))


async def main(test_dir):
    env = {**os.environ,
           "TRACKER_FILE_OVERRIDE": os.path.join(test_dir, "List.xlsx"),
           "HOT_JOBS_FILE_OVERRIDE": os.path.join(test_dir, "daily_hot_jobs.json"),
           "REJECTED_REMOTE_FILE_OVERRIDE": os.path.join(test_dir, "rejected_remote.json"),
           "PREVIOUS_JOBS_FILE_OVERRIDE": os.path.join(test_dir, "previous_jobs.json")}
    hot = lambda: json.load(open(env["HOT_JOBS_FILE_OVERRIDE"], encoding="utf-8"))
    remote_count = lambda: len(json.load(open(env["REJECTED_REMOTE_FILE_OVERRIDE"], encoding="utf-8")))
    params = StdioServerParameters(command=sys.executable, args=[os.path.join(HERE, "server.py")], env=env)
    passed = failed = 0

    def check(label, cond, detail=""):
        nonlocal passed, failed
        print(("PASS  " if cond else "FAIL  ") + label + ("" if cond else f"  -> {detail}"))
        passed += cond
        failed += not cond

    async with stdio_client(params) as (read, write):
        async with ClientSession(read, write) as s:
            await s.initialize()
            tools = {t.name: t for t in (await s.list_tools()).tools}
            print("tools:", ", ".join(tools))
            check("9 tools exposed", set(tools) == {"search_tracker", "add_job", "mark_rejected", "add_hr_contact",
                                                     "block_hot_job", "block_all_hot_jobs", "reject_remote_job",
                                                     "reject_all_remote", "update_status"}, tools)
            check("add_job schema has company + status", {"company", "status"} <= set(tools["add_job"].inputSchema["properties"]))

            async def call(tool_name, **kw):
                r = await s.call_tool(tool_name, kw)
                if r.isError:
                    return {"ok": False, "outcome": "tool_error", "message": r.content[0].text}
                return json.loads(r.content[0].text)

            r = await call("search_tracker", query="Bluecoders")
            check("search_tracker finds Bluecoders rows", r["ok"] and len(r["results"]) >= 5, r)
            r = await call("search_tracker", query="https://jobs.lever.co/yuno/09574e89-949a-451a-9ef6-efaa62d96b31")
            check("search_tracker matches URL (row 243)", any(x["row"] == 243 for x in r["results"]), r)

            r = await call("add_job", company="MCP Test Co", role="Backend", url="https://example.com/mcp-1")
            check("add_job adds a new job", r["outcome"] == "added", r)
            r = await call("add_job", company="MCP Test Co", role="Backend", url="https://example.com/mcp-1/thanks")
            check("add_job refuses the same job twice", r["outcome"] == "duplicate", r)
            r = await call("add_job", company="MCP Test Co", role="Other", url="https://example.com/mcp-2")
            check("add_job flags existing company", r["outcome"] == "company_exists", r)
            r = await call("add_job", company="X", status="applied")
            check("add_job rejects invalid status", r["outcome"] == "error", r)

            r = await call("mark_rejected", company="MCP Test Co")
            check("mark_rejected previews without confirm", r["outcome"] == "preview" and "DRY RUN" in r["message"], r)
            r = await call("search_tracker", query="MCP Test Co")
            check("...and nothing changed", r["results"][0]["status"] == "done", r)
            r = await call("mark_rejected", company="MCP Test Co", confirm=True)
            check("mark_rejected with confirm updates", r["outcome"] == "rejected", r)

            # --- update_status ---
            r = await call("update_status", company="MCP Test Co", status="In progress", note="phone screen Monday")
            check("update_status single row", r["outcome"] == "updated", r)
            import openpyxl
            ws = openpyxl.load_workbook(env["TRACKER_FILE_OVERRIDE"])["Sheet1"]
            row = next(i for i in range(2, ws.max_row + 1) if ws.cell(i, 1).value == "MCP Test Co")
            check("...status In progress, strikethrough removed, dated note saved",
                  ws.cell(row, 4).value == "In progress" and not ws.cell(row, 1).font.strike
                  and "phone screen Monday" in (ws.cell(row, 6).value or ""), (ws.cell(row, 4).value, ws.cell(row, 6).value))
            r = await call("update_status", company="Galadrim", status="In progress")
            check("update_status refuses ambiguous company (2 Galadrim rows)", r["outcome"] == "ambiguous"
                  and "Row 352" in r["message"] and "Row 353" in r["message"], r)
            r = await call("update_status", row=352, status="In progress")
            check("update_status by row", r["outcome"] == "updated" and "Row 352" in r["message"], r)
            r = await call("update_status", company="MCP Test Co", status="applied")
            check("update_status rejects invalid status", r["outcome"] == "error", r)

            r = await call("add_hr_contact", company="MCP Test Co", name="Jane Recruiter", url="https://www.linkedin.com/in/jane", title="TA")
            check("add_hr_contact appends", r["outcome"] == "added", r)
            r = await call("add_hr_contact", company="MCP Test Co", name="Jane Recruiter")
            check("add_hr_contact doesn't duplicate", r["outcome"] == "unchanged", r)

            # --- daily digest hot jobs (seeded into the copy) ---
            data = hot()
            data["current_jobs"]["Backend Java"] = [
                {"company": "Hot Test A", "title": "Java Dev", "url": "https://example.com/a"},
                {"company": "Hot Test B", "title": "Tech Lead", "url": "https://example.com/b"},
                {"company": "MCP Test Co", "title": "Backend", "url": "https://example.com/c"},  # in tracker
            ]
            json.dump(data, open(env["HOT_JOBS_FILE_OVERRIDE"], "w", encoding="utf-8"), ensure_ascii=False)
            bl_before = len(hot()["blocklist"])

            r = await call("block_hot_job", company="Hot Test A", title="Java Dev")
            titles = [j["company"] for j in hot()["current_jobs"]["Backend Java"]]
            check("block_hot_job removes it and blocklists", r["ok"] and "Hot Test A" not in titles
                  and len(hot()["blocklist"]) == bl_before + 1, r)
            r = await call("block_all_hot_jobs")
            check("block_all_hot_jobs previews, skips tracker company", r["outcome"] == "preview"
                  and "Hot Test B" in r["message"] and "Skipped" in r["message"]
                  and len(hot()["current_jobs"]["Backend Java"]) == 2, r)
            r = await call("block_all_hot_jobs", confirm=True)
            left = [j["company"] for j in hot()["current_jobs"]["Backend Java"]]
            check("block_all_hot_jobs confirm blocks all except tracker company", left == ["MCP Test Co"], left)

            # --- remote digest ---
            json.dump([["Remote Test Co", "Senior Backend Engineer"], ["Remote Other", "Java Dev"]],
                      open(env["PREVIOUS_JOBS_FILE_OVERRIDE"], "w", encoding="utf-8"))
            n0 = remote_count()
            r = await call("reject_all_remote")
            check("reject_all_remote previews last digest", r["outcome"] == "preview"
                  and "Remote Test Co" in r["message"] and remote_count() == n0, r)
            r = await call("reject_remote_job", company="Remote Only Co", title="")
            check("reject_remote_job adds entry", r["ok"] and remote_count() == n0 + 1, r)
            r = await call("reject_all_remote", confirm=True)
            check("reject_all_remote confirm adds last digest", r["ok"] and remote_count() == n0 + 3, r)

    print(f"RESULT: {passed} passed, {failed} failed")
    return failed


if __name__ == "__main__":
    d = sys.argv[1]
    missing = [f for f in ("List.xlsx", "daily_hot_jobs.json", "rejected_remote.json", "previous_jobs.json")
               if not os.path.exists(os.path.join(d, f))]
    if missing:
        sys.exit(f"Test dir is missing copies of: {missing}")
    sys.exit(asyncio.run(main(d)))
