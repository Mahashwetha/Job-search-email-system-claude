"""
tracker_lib.py - shared, safe access to List.xlsx for the tracker scripts.

Used by:
  .claude/skills/search/scripts/search_tracker.py
  .claude/skills/new-job/scripts/add_job.py
  .claude/skills/update-hr/scripts/add_hr_contact.py

Set TRACKER_FILE_OVERRIDE to point every script at a copy (used for tests).
"""
import os
import re
import shutil
import sys
import tempfile
import time
from contextlib import contextmanager
from datetime import datetime

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from config import TRACKER_FILE as _CONFIG_TRACKER  # noqa: E402

TRACKER_FILE = os.environ.get("TRACKER_FILE_OVERRIDE") or _CONFIG_TRACKER
BACKUP_DIR = os.path.join(os.path.dirname(TRACKER_FILE), "backups")
LOCK_FILE = TRACKER_FILE + ".lock"
SHEETS = ("Sheet1",)  # the old "Rejected" sheet was merged into Sheet1 on 2026-09-25
COL_COMPANY, COL_ROLE, COL_URL, COL_STATUS, COL_HR, COL_NOTES = 1, 2, 3, 4, 5, 6
VALID_STATUSES = {"done", "In progress", "Rejected"}


@contextmanager
def tracker_lock(timeout=30, stale_after=300):
    """One writer at a time (Claude skills, email agent, MCP tools all share this)."""
    start = time.time()
    while True:
        try:
            fd = os.open(LOCK_FILE, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
            os.write(fd, f"{os.getpid()} {datetime.now().isoformat()}".encode())
            os.close(fd)
            break
        except FileExistsError:
            if time.time() - os.path.getmtime(LOCK_FILE) > stale_after:
                os.remove(LOCK_FILE)  # left behind by a crashed run
                continue
            if time.time() - start > timeout:
                raise TimeoutError(f"Tracker is busy (lock held: {LOCK_FILE})")
            time.sleep(0.5)
    try:
        yield
    finally:
        if os.path.exists(LOCK_FILE):
            os.remove(LOCK_FILE)


def backup():
    os.makedirs(BACKUP_DIR, exist_ok=True)
    path = os.path.join(BACKUP_DIR, f"List_backup_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
    shutil.copy2(TRACKER_FILE, path)
    return path


def load(data_only=False):
    """Open the workbook; if Excel has it locked, read from a temp copy."""
    try:
        return openpyxl.load_workbook(TRACKER_FILE, data_only=data_only)
    except PermissionError:
        fd, tmp = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        shutil.copy2(TRACKER_FILE, tmp)
        try:
            return openpyxl.load_workbook(tmp, data_only=data_only)
        finally:
            os.remove(tmp)


def save(wb):
    try:
        wb.save(TRACKER_FILE)
    except PermissionError:
        raise PermissionError("List.xlsx is open in Excel. Close it and run again (nothing was saved).")


def cell_text(value):
    return "" if value is None else str(value).strip()


def extract_url(value):
    """Plain URL, or the URL inside an =HYPERLINK("url","text") formula."""
    text = cell_text(value)
    m = re.search(r'HYPERLINK\(\s*"([^"]+)"', text, re.I)
    return m.group(1) if m else text


def norm_url(url):
    """Compare job links ignoring scheme, www, query string, trailing slash and /thanks, /application."""
    u = cell_text(url).lower()
    if not u.startswith(("http://", "https://")):
        return u
    u = re.sub(r"^https?://(www\.)?", "", u)
    u = u.split("?")[0].split("#")[0].rstrip("/")
    u = re.sub(r"/(thanks|application|apply)$", "", u)
    return u


def rows(wb):
    """Yield (sheet_name, row_number, [A..F values]) for every data row."""
    for name in SHEETS:
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        for r in range(2, ws.max_row + 1):
            vals = [ws.cell(r, c).value for c in range(1, 7)]
            if any(cell_text(v) for v in vals):
                yield name, r, vals


def company_matches(query, company):
    q, c = query.strip().lower(), cell_text(company).lower()
    return bool(q and c) and (q in c or c in q)
