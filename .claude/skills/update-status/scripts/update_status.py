"""
update_status.py - change the status of an existing tracker row (e.g. a callback -> 'In progress').

Usage:
  python .claude/skills/update-status/scripts/update_status.py --company "Galadrim" --status "In progress" [--note "phone screen booked"]
  python .claude/skills/update-status/scripts/update_status.py --row 352 --status "In progress"
  add --all to update every matching row, --dry-run to preview

Rules enforced in code:
  - status must be 'done', 'In progress' or 'Rejected'
  - company matching several rows -> refused unless --row or --all (lists the rows)
  - 'Rejected' adds strikethrough; any other status removes it
  - --note is appended to column F as "YYYY-MM-DD: note" (existing comments kept)
  - shared tracker lock + backup before writing
Exit codes: 0 updated / previewed, 4 no match, 5 ambiguous (several rows), 1 error.
"""
import argparse
import os
import sys
from datetime import date

from openpyxl.styles import Font

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "..", ".."))
import tracker_lib as T  # noqa: E402


def set_strike(ws, row, strike):
    for col in range(1, 7):
        cell = ws.cell(row=row, column=col)
        f = cell.font
        cell.font = Font(name=f.name, size=f.size, bold=f.bold, italic=f.italic,
                         color=f.color, underline=f.underline, strike=strike)


def main():
    p = argparse.ArgumentParser()
    g = p.add_mutually_exclusive_group(required=True)
    g.add_argument("--company")
    g.add_argument("--row", type=int)
    p.add_argument("--status", required=True)
    p.add_argument("--note", default="")
    p.add_argument("--all", action="store_true", help="update every row matching --company")
    p.add_argument("--dry-run", action="store_true")
    a = p.parse_args()

    if a.status not in T.VALID_STATUSES:
        print(f"ERROR: status must be one of {sorted(T.VALID_STATUSES)} (got '{a.status}')")
        return 1

    with T.tracker_lock():
        wb = T.load()
        ws = wb["Sheet1"]
        if a.row:
            if a.row < 2 or a.row > ws.max_row or not T.cell_text(ws.cell(a.row, 1).value):
                print(f"ERROR: row {a.row} is empty or outside the tracker (2..{ws.max_row})")
                return 1
            targets = [a.row]
        else:
            q = a.company.strip().lower()
            targets = [r for _, r, v in T.rows(wb) if q and q in T.cell_text(v[0]).lower()]
            if not targets:
                print(f"No rows found matching: '{a.company}'")
                return 4
            if len(targets) > 1 and not a.all:
                print(f"AMBIGUOUS: '{a.company}' matches {len(targets)} rows. Re-run with --row N (or --all):")
                for r in targets:
                    print(f"  Row {r}: {T.cell_text(ws.cell(r, 1).value)} | {T.cell_text(ws.cell(r, 2).value)} | {T.cell_text(ws.cell(r, 4).value)}")
                return 5

        changes = [(r, T.cell_text(ws.cell(r, 1).value), T.cell_text(ws.cell(r, 2).value),
                    T.cell_text(ws.cell(r, 4).value)) for r in targets]
        if a.dry_run:
            print(f"DRY RUN: would set status '{a.status}' on {len(changes)} row(s):")
            for r, co, role, old in changes:
                print(f"  Row {r}: {co} | {role} | {old or '(empty)'} -> {a.status}")
            return 0

        backup_path = T.backup()
        for r, *_ in changes:
            ws.cell(r, T.COL_STATUS).value = a.status
            set_strike(ws, r, a.status == "Rejected")
            if a.note.strip():
                old = T.cell_text(ws.cell(r, T.COL_NOTES).value)
                entry = f"{date.today():%Y-%m-%d}: {a.note.strip()}"
                ws.cell(r, T.COL_NOTES).value = f"{old}\n{entry}" if old else entry
        T.save(wb)
    for r, co, role, old in changes:
        print(f"UPDATED: Row {r}: {co} | {role} | {old or '(empty)'} -> {a.status}")
    print(f"(backup: {backup_path})")
    return 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    try:
        sys.exit(main())
    except (TimeoutError, PermissionError) as e:
        print(f"ERROR: {e}")
        sys.exit(1)
