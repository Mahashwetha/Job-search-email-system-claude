"""
add_hr_contact.py - append ONE HR contact to column E of a tracker row. Never overwrites.

Usage:
  python .claude/skills/update-hr/scripts/add_hr_contact.py --row 311 --name "Jane Doe" --url "https://www.linkedin.com/in/..." [--title "Tech Recruiter"] [--email "x@y.com"]
  python .claude/skills/update-hr/scripts/add_hr_contact.py --company "Bluecoders" --name ...   (all rows of that company)

Rules enforced in code:
  - existing contacts in the cell are always kept; the new one is appended with CHAR(10)
  - if the name is already in the cell, nothing changes (safe to re-run)
  - same HYPERLINK formula format as update_hr_contacts.py
  - takes the shared tracker lock and backs up List.xlsx before writing
For bulk updates from the HR_CONTACTS dict, keep using update_hr_contacts.py.
"""
import argparse
import os
import sys

from openpyxl.styles import Alignment, Font

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "..", ".."))
import tracker_lib as T  # noqa: E402


def contact_part(name, url, title, email):
    label = name
    if title:
        label += f" ({title})"
    if email:
        label += f" {email}"
    label = label.replace('"', "'")
    return f'HYPERLINK("{url}","{label}")' if url else f'"{label}"'


def main():
    p = argparse.ArgumentParser()
    g = p.add_mutually_exclusive_group(required=True)
    g.add_argument("--row", type=int)
    g.add_argument("--company")
    p.add_argument("--name", required=True)
    p.add_argument("--url", default="")
    p.add_argument("--title", default="")
    p.add_argument("--email", default="")
    a = p.parse_args()

    with T.tracker_lock():
        wb = T.load()
        ws = wb["Sheet1"]
        if a.row:
            if a.row < 2 or a.row > ws.max_row:
                print(f"ERROR: row {a.row} is outside the tracker (2..{ws.max_row})")
                return 1
            targets = [a.row]
        else:
            targets = [r for _, r, v in T.rows(wb) if T.company_matches(a.company, v[0])]
            if not targets:
                print(f"ERROR: company '{a.company}' not in tracker. Add the job first.")
                return 1

        part = contact_part(a.name.strip(), a.url.strip(), a.title.strip(), a.email.strip())
        changed, skipped = [], []
        for r in targets:
            cell = ws.cell(r, T.COL_HR)
            existing = T.cell_text(cell.value)
            if a.name.strip().lower() in existing.lower():
                skipped.append(r)
                continue
            if not existing:
                formula = "=" + part if part.startswith("HYPERLINK") else part.strip('"')
            elif existing.startswith("="):
                formula = existing + " & CHAR(10) & " + part
            else:
                formula = "=" + '"' + existing.replace('"', "'") + '"' + " & CHAR(10) & " + part
            cell.value = formula
            cell.number_format = "General"
            cell.font = Font(color="0563C1", underline="single", size=10, strike=cell.font.strike)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            lines = formula.count("CHAR(10)") + 1
            ws.row_dimensions[r].height = max(30, lines * 15)
            changed.append(r)

        if not changed:
            print(f"NO CHANGE: '{a.name}' already listed on row(s) {skipped}")
            return 0
        backup_path = T.backup()
        T.save(wb)
        print(f"ADDED CONTACT: {a.name} -> row(s) {changed}" + (f" (already on {skipped})" if skipped else ""))
        print(f"(backup: {backup_path})")
        return 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    try:
        sys.exit(main())
    except (TimeoutError, PermissionError) as e:
        print(f"ERROR: {e}")
        sys.exit(1)
